from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from typing import Any

import pytest
from openpyxl import load_workbook

from config import today
from enums import LedgerKind
from models import LedgerEntry
from schemas import format_invoice_date
from tests.helpers import InvoiceApi


def shifted(days: int) -> date:
    return today() + timedelta(
        days=days,
    )


def day(shift: int = 0) -> str:
    """a date relative to today in the api format"""
    return format_invoice_date(
        value=shifted(
            days=shift,
        ),
    )


class LedgerApi(InvoiceApi):
    def stock(
        self,
        amount: str,
        shift: int = 0,
    ) -> LedgerEntry:
        """an invoice inflow without the invoice, the fake ledger does not need one"""
        entry = LedgerEntry(
            id=self.store.new_id(),
            kind=LedgerKind.INVOICE,
            amount=Decimal(
                value=amount,
            ),
            cash=Decimal(
                value=0,
            ),
            occurred_on=shifted(
                days=shift,
            ),
        )
        self.store.entries.append(
            entry,
        )
        return entry

    def post(
        self,
        url: str,
        payload: dict[str, Any] | None = None,
        expected: int = 200,
    ) -> dict[str, Any]:
        response = self.client.post(
            url=url,
            json=payload,
        )
        assert response.status_code == expected, response.text
        return response.json()

    def sale(
        self,
        amount: str,
        cash: str = "0.00",
        shift: int = 0,
        expected: int = 200,
    ) -> dict[str, Any]:
        """the whole sale, the cash part goes into the till and the rest is cashless"""
        card = Decimal(
            value=amount,
        ) - Decimal(
            value=cash,
        )
        return self.post(
            url="/ledger/sale",
            payload={
                "occurred_on": day(
                    shift=shift,
                ),
                "cash": cash,
                "card": str(
                    card,
                ),
                "comment": None,
            },
            expected=expected,
        )

    def cash(
        self,
        kind: str,
        amount: str,
        comment: str | None = None,
        shift: int = 0,
        expected: int = 200,
    ) -> dict[str, Any]:
        """a cash-in or a cash-out"""
        return self.post(
            url=f"/ledger/{kind}",
            payload={
                "occurred_on": day(
                    shift=shift,
                ),
                "amount": amount,
                "comment": comment,
            },
            expected=expected,
        )

    def cash_count(
        self,
        counted: str,
        shift: int = 0,
        expected: int = 200,
    ) -> dict[str, Any]:
        return self.post(
            url="/ledger/cash-count",
            payload={
                "occurred_on": day(
                    shift=shift,
                ),
                "counted": counted,
                "comment": None,
            },
            expected=expected,
        )

    def stocktake(
        self,
        counted: str,
        shift: int = 0,
    ) -> dict[str, Any]:
        return self.post(
            url="/ledger/stocktake",
            payload={
                "occurred_on": day(
                    shift=shift,
                ),
                "counted": counted,
                "comment": "переучёт",
            },
        )

    def reverse(
        self,
        entry_id: int,
        expected: int = 200,
    ) -> dict[str, Any]:
        return self.post(
            url=f"/ledger/{entry_id}/reverse",
            expected=expected,
        )

    def history(self) -> dict[str, Any]:
        return self.client.get(
            url="/ledger",
        ).json()

    def revalue(
        self,
        lines: list[tuple[int, str, int]],
        shift: int = 0,
        expected: int = 200,
    ) -> dict[str, Any]:
        """product ids with the counted stock and the new price"""
        return self.post(
            url="/ledger/revaluation",
            payload={
                "occurred_on": day(
                    shift=shift,
                ),
                "comment": "акция",
                "lines": [
                    {
                        "product_id": product_id,
                        "quantity": quantity,
                        "new_price": new_price,
                    }
                    for product_id, quantity, new_price in lines
                ],
            },
            expected=expected,
        )


class TestOutflow(LedgerApi):
    def test_sale_and_write_off_take_stock_out(self) -> None:
        self.stock(
            amount="10000.00",
        )

        sale = self.sale(
            amount="1500.00",
            shift=-1,
        )
        write_off = self.post(
            url="/ledger/write-off",
            payload={
                "occurred_on": day(),
                "amount": "200.00",
                "comment": "забрал владелец",
            },
        )

        assert (sale["amount"], sale["balance"]) == ("-1500.00", "8500.00")
        assert (write_off["amount"], write_off["balance"]) == ("-200.00", "8300.00")
        entry = self.store.entries[-1]
        assert (entry.kind, entry.comment) == (LedgerKind.WRITE_OFF, "забрал владелец")
        assert self.store.locks == 2

    @pytest.mark.parametrize(
        argnames=("amount", "cash", "shift", "detail"),
        argvalues=[
            ("0", "0.00", 0, "Укажите выручку наличными или безналом"),
            ("10.00", "0.00", 1, "Дата не может быть в будущем"),
            # more cash than the whole sale leaves a negative cashless part
            ("10.00", "15.00", 0, ["body", "card"]),
        ],
    )
    def test_returns_422_when_sale_amount_or_date_is_invalid(
        self,
        amount: str,
        cash: str,
        shift: int,
        detail: str | list[str],
    ) -> None:
        response = self.sale(
            amount=amount,
            cash=cash,
            shift=shift,
            expected=422,
        )

        if isinstance(
            detail,
            str,
        ):
            assert response == {"detail": detail}
        else:
            assert [error["loc"] for error in response["detail"]] == [detail]
        assert self.store.entries == []


class TestCash(LedgerApi):
    def test_the_till_follows_the_takings_and_the_payments(self) -> None:
        self.stock(
            amount="200000.00",
        )

        sale = self.sale(
            amount="40000.00",
            cash="15000.00",
        )
        taken = self.cash(
            kind="cash-out",
            amount="10000.00",
            comment="забрал владелец",
        )
        put = self.cash(
            kind="cash-in",
            amount="20000.00",
        )
        paid = self.cash(
            kind="cash-out",
            amount="3200.00",
            comment="накладная № 17",
        )

        # the cashless part is gone to the bank, only the cash stays in the till
        assert (
            sale["amount"],
            sale["balance"],
            sale["cash"],
            sale["cash_balance"],
        ) == (
            "-40000.00",
            "160000.00",
            "15000.00",
            "15000.00",
        )
        assert [
            (posted["cash"], posted["cash_balance"], posted["balance"])
            for posted in (taken, put, paid)
        ] == [
            ("-10000.00", "5000.00", "160000.00"),
            ("20000.00", "25000.00", "160000.00"),
            ("-3200.00", "21800.00", "160000.00"),
        ]
        page = self.history()
        assert (page["balance"], page["cash_balance"]) == ("160000.00", "21800.00")
        assert [
            (row["kind"], row["amount"], row["cash"], row["cash_after"])
            for row in page["entries"][:2]
        ] == [
            ("cash_out", "0", "-3200.00", "21800.00"),
            ("cash_in", "0", "20000.00", "25000.00"),
        ]
        assert self.store.entries[-1].comment == "накладная № 17"

    def test_cash_out_needs_a_comment(self) -> None:
        response = self.cash(
            kind="cash-out",
            amount="100.00",
            expected=422,
        )

        assert [error["loc"] for error in response["detail"]] == [["body", "comment"]]
        assert self.store.entries == []

    def test_cash_count_records_the_shortage(self) -> None:
        self.cash(
            kind="cash-in",
            amount="5000.00",
            shift=-1,
        )
        self.sale(
            amount="1000.00",
            cash="600.00",
        )

        posted = self.cash_count(
            counted="5500.00",
        )

        assert (posted["amount"], posted["cash"], posted["cash_balance"]) == (
            "0",
            "-100.00",
            "5500.00",
        )
        balance = self.client.get(
            url="/ledger/balance",
        ).json()
        assert (
            balance["cash_balance"],
            balance["cash_closed_on"],
            balance["closed_on"],
        ) == (
            "5500.00",
            day(),
            None,
        )

    def test_a_count_closes_only_what_it_counted(self) -> None:
        self.stock(
            amount="1000.00",
            shift=-3,
        )
        self.stocktake(
            counted="1000.00",
        )
        self.cash_count(
            counted="0",
        )

        # the stock and the cash both closed today
        assert (
            self.post(
                url="/ledger/write-off",
                payload={
                    "occurred_on": day(
                        shift=-1,
                    ),
                    "amount": "10.00",
                    "comment": None,
                },
                expected=409,
            )["detail"]
            == f"Период закрыт переучётом от {day()}"
        )
        assert (
            self.cash(
                kind="cash-in",
                amount="10.00",
                shift=-1,
                expected=409,
            )["detail"]
            == f"Касса закрыта сверкой от {day()}"
        )

    def test_a_cash_count_leaves_the_stock_open(self) -> None:
        self.stock(
            amount="1000.00",
            shift=-3,
        )
        self.cash_count(
            counted="0",
        )

        self.post(
            url="/ledger/write-off",
            payload={
                "occurred_on": day(
                    shift=-1,
                ),
                "amount": "10.00",
                "comment": None,
            },
        )
        # a cashless sale does not touch the till
        self.sale(
            amount="20.00",
            shift=-1,
        )
        blocked = self.sale(
            amount="20.00",
            cash="5.00",
            shift=-1,
            expected=409,
        )

        assert blocked["detail"] == f"Касса закрыта сверкой от {day()}"

    def test_a_stocktake_leaves_the_till_open(self) -> None:
        self.stock(
            amount="1000.00",
            shift=-3,
        )
        self.stocktake(
            counted="1000.00",
        )

        posted = self.cash(
            kind="cash-in",
            amount="300.00",
            shift=-1,
        )

        assert posted["cash_balance"] == "300.00"

    def test_reversal_of_a_sale_returns_the_stock_and_the_cash(self) -> None:
        self.stock(
            amount="1000.00",
        )
        sale = self.sale(
            amount="400.00",
            cash="150.00",
        )

        reversal = self.reverse(
            entry_id=sale["id"],
        )

        assert (reversal["amount"], reversal["cash"]) == ("400.00", "-150.00")
        assert (reversal["balance"], reversal["cash_balance"]) == ("1000.00", "0.00")

    def test_cash_count_blocks_earlier_cash_reversals(self) -> None:
        put = self.cash(
            kind="cash-in",
            amount="500.00",
            shift=-1,
        )
        sale = self.sale(
            amount="100.00",
            shift=-1,
        )
        self.cash_count(
            counted="500.00",
        )

        blocked = self.reverse(
            entry_id=put["id"],
            expected=409,
        )
        # a cashless sale changed nothing the count saw in the till
        self.reverse(
            entry_id=sale["id"],
        )

        assert (
            blocked["detail"]
            == "После этой записи была сверка кассы, исправьте кассу новой сверкой"
        )


class TestStocktake(LedgerApi):
    def test_entry_is_the_difference_and_the_balance_becomes_the_count(self) -> None:
        self.stock(
            amount="10000.00",
        )
        self.sale(
            amount="1000.00",
        )

        posted = self.stocktake(
            counted="8700.00",
        )

        assert (posted["amount"], posted["balance"]) == ("-300.00", "8700.00")
        [latest, *_] = self.history()["entries"]
        assert (latest["kind"], latest["amount"], latest["balance_after"]) == (
            "stocktake",
            "-300.00",
            "8700.00",
        )

    def test_counts_against_the_balance_on_its_date(self) -> None:
        self.stock(
            amount="5000.00",
            shift=-3,
        )
        self.sale(
            amount="400.00",
        )

        posted = self.stocktake(
            counted="5100.00",
            shift=-1,
        )

        # the sale of today comes after the stocktake of yesterday
        assert (posted["amount"], posted["balance"]) == ("100.00", "4700.00")

    def test_stocktake_closes_the_period(self) -> None:
        self.stock(
            amount="1000.00",
            shift=-3,
        )
        self.stocktake(
            counted="900.00",
        )

        blocked = self.sale(
            amount="10.00",
            shift=-1,
            expected=409,
        )
        self.sale(
            amount="10.00",
        )

        assert blocked["detail"] == f"Период закрыт переучётом от {day()}"
        assert self.client.get(
            url="/ledger/balance",
            params={
                "on": day(
                    shift=-1,
                ),
            },
        ).json() == {
            "balance": "1000.00",
            "cash_balance": "0",
            "closed_on": day(),
            "cash_closed_on": None,
        }


class TestReverse(LedgerApi):
    def test_reversal_cancels_sale_and_links_original_entry(self) -> None:
        self.stock(
            amount="1000.00",
        )
        sale = self.sale(
            amount="250.00",
            shift=-2,
        )

        reversal = self.reverse(
            entry_id=sale["id"],
        )

        assert reversal == {
            "id": self.store.entries[-1].id,
            "amount": "250.00",
            "cash": "0.00",
            "balance": "1000.00",
            "cash_balance": "0.00",
        }
        entry = self.store.entries[-1]
        assert (entry.kind, entry.reverses_id, entry.comment, entry.occurred_on) == (
            LedgerKind.SALE,
            sale["id"],
            None,
            shifted(
                days=-2,
            ),
        )
        [cancelled] = [
            row for row in self.history()["entries"] if row["id"] == sale["id"]
        ]
        assert cancelled["reversed_by_id"] == reversal["id"]

    @pytest.mark.parametrize(
        argnames=("target", "detail"),
        argvalues=[
            ("sale", "Запись уже сторнирована"),
            ("reversal", "Сторно нельзя сторнировать"),
        ],
    )
    def test_returns_409_when_entry_cannot_be_reversed_again(
        self,
        target: str,
        detail: str,
    ) -> None:
        sale = self.sale(
            amount="250.00",
        )
        reversal = self.reverse(
            entry_id=sale["id"],
        )
        entry_id = sale["id"] if target == "sale" else reversal["id"]

        response = self.reverse(
            entry_id=entry_id,
            expected=409,
        )

        assert response == {"detail": detail}

    def test_invoice_and_unknown_entries_are_not_reversed(self) -> None:
        inflow = self.stock(
            amount="1000.00",
        )

        assert (
            self.reverse(
                entry_id=inflow.id,
                expected=409,
            )["detail"]
            == "Проведённую накладную нельзя сторнировать"
        )
        self.reverse(
            entry_id=999,
            expected=404,
        )

    def test_stocktake_blocks_earlier_reversals_until_it_is_reversed(self) -> None:
        self.stock(
            amount="1000.00",
        )
        sale = self.sale(
            amount="100.00",
            shift=-1,
        )
        stocktake = self.stocktake(
            counted="850.00",
        )

        blocked = self.reverse(
            entry_id=sale["id"],
            expected=409,
        )
        self.reverse(
            entry_id=stocktake["id"],
        )
        reversal = self.reverse(
            entry_id=sale["id"],
        )

        assert (
            blocked["detail"]
            == "После этой записи был переучёт, исправьте остаток новым переучётом"
        )
        assert reversal["balance"] == "1000.00"


class TestRevaluation(LedgerApi):
    def test_reprices_the_stock_on_the_shelf(self) -> None:
        self.stock(
            amount="1000.00",
        )
        _, sugar, salt = self.known_supplier()

        posted = self.revalue(
            lines=[
                (sugar.id, "4", 60),
                (salt.id, "2.5", 40),
            ],
        )

        # sugar 4 × (60 − 65) and salt 2.5 × (40 − 36)
        assert (posted["amount"], posted["balance"]) == ("-10.00", "990.00")
        entry = self.store.entries[-1]
        assert (entry.kind, entry.comment) == (LedgerKind.REVALUATION, "акция")
        assert [
            (item.product_id, item.old_price, item.new_price, item.invoice_item_id)
            for item in self.store.revaluations
        ] == [
            (sugar.id, 65, 60, None),
            (salt.id, 36, 40, None),
        ]
        assert (sugar.retail_price, salt.retail_price, sugar.markup_pct) == (60, 40, 30)

    @pytest.mark.parametrize(
        argnames=("lines", "detail"),
        argvalues=[
            ([(999, "1", 10)], "Товара с id 999 нет"),
            (
                [(2, "1", 70), (2, "2", 70)],
                "Товар «Сахар-песок 1 кг» указан дважды",
            ),
            ([(2, "1", 65)], "Цена товара «Сахар-песок 1 кг» уже 65 ₽"),
        ],
    )
    def test_rejects_unknown_repeated_and_unchanged_products(
        self,
        lines: list[tuple[int, str, int]],
        detail: str,
    ) -> None:
        _, sugar, _ = self.known_supplier()

        assert (
            self.revalue(
                lines=lines,
                expected=422,
            )["detail"]
            == detail
        )

        assert (self.store.entries, sugar.retail_price) == ([], 65)

    def test_closed_period_is_not_revalued(self) -> None:
        _, sugar, _ = self.known_supplier()
        self.stocktake(
            counted="0",
        )

        self.revalue(
            lines=[
                (sugar.id, "1", 70),
            ],
            shift=-1,
            expected=409,
        )

    def test_reversal_returns_the_old_price(self) -> None:
        _, sugar, salt = self.known_supplier()
        posted = self.revalue(
            lines=[
                (sugar.id, "4", 70),
                (salt.id, "1", 40),
            ],
        )
        # the salt got another price since, it is not rolled back
        salt.retail_price = Decimal(
            value=45,
        )

        reversal = self.reverse(
            entry_id=posted["id"],
        )

        assert (reversal["amount"], reversal["balance"]) == ("-24.00", "0.00")
        assert (sugar.retail_price, salt.retail_price) == (65, 45)


class TestProductEdit(LedgerApi):
    def edit(
        self,
        product_id: int,
        name: str,
        retail_price: int,
        quantity: str | None = None,
        shift: int = 0,
        expected: int = 200,
    ) -> Any:
        """without a quantity the price changes without a revaluation"""
        response = self.client.put(
            url=f"/product/{product_id}",
            json={
                "name": name,
                "retail_price": retail_price,
                "revaluation": None
                if quantity is None
                else {
                    "occurred_on": day(
                        shift=shift,
                    ),
                    "quantity": quantity,
                    "comment": None,
                },
            },
        )
        assert response.status_code == expected, response.text
        return response.json()

    def spellings(self, product_id: int) -> list[str]:
        return [
            alias.alias_raw
            for alias in self.store.aliases
            if alias.product_id == product_id
        ]

    def test_corrected_name_keeps_the_old_one_as_a_spelling(self) -> None:
        _, sugar, _ = self.known_supplier()

        assert (
            self.edit(
                product_id=sugar.id,
                name="  Сахар-песок 1 кг ГОСТ ",
                retail_price=65,
            )
            is None
        )

        assert sugar.name == "Сахар-песок 1 кг ГОСТ"
        assert self.spellings(
            product_id=sugar.id,
        ) == ["Сахар-песок 1 кг", "Сахар-песок 1 кг ГОСТ"]
        assert self.store.entries == []
        catalog = self.client.get(
            url="/catalog",
        ).json()
        assert [
            product["spellings"]
            for product in catalog["products"]
            if product["id"] == sugar.id
        ] == [["Сахар-песок 1 кг"]]

    def test_a_case_fix_adds_no_spelling(self) -> None:
        _, sugar, _ = self.known_supplier()

        self.edit(
            product_id=sugar.id,
            name="САХАР-песок 1 кг",
            retail_price=65,
        )

        assert sugar.name == "САХАР-песок 1 кг"
        assert self.spellings(
            product_id=sugar.id,
        ) == ["Сахар-песок 1 кг"]

    def test_name_of_another_product_is_rejected(self) -> None:
        _, sugar, _ = self.known_supplier()

        assert self.edit(
            product_id=sugar.id,
            name="«Соль»",
            retail_price=70,
            expected=409,
        ) == {"detail": "Название ««Соль»» уже у товара «Соль»"}

        assert (sugar.name, sugar.retail_price) == ("Сахар-песок 1 кг", 65)
        assert self.spellings(
            product_id=sugar.id,
        ) == ["Сахар-песок 1 кг"]

    def test_price_without_old_stock_changes_without_an_entry(self) -> None:
        _, sugar, _ = self.known_supplier()

        self.edit(
            product_id=sugar.id,
            name="Сахар-песок 1 кг",
            retail_price=70,
        )

        assert (sugar.retail_price, sugar.markup_pct) == (70, 30)
        assert (self.store.entries, self.store.revaluations) == ([], [])

    def test_price_with_old_stock_revalues_it(self) -> None:
        self.stock(
            amount="1000.00",
        )
        _, sugar, _ = self.known_supplier()

        posted = self.edit(
            product_id=sugar.id,
            name="Сахар 1 кг",
            retail_price=70,
            quantity="4",
        )

        # 4 × (70 − 65)
        assert (posted["amount"], posted["balance"]) == ("20.00", "1020.00")
        assert self.store.entries[-1].kind == LedgerKind.REVALUATION
        assert [
            (item.product_id, item.quantity, item.old_price, item.new_price)
            for item in self.store.revaluations
        ] == [(sugar.id, 4, 65, 70)]
        assert (sugar.name, sugar.retail_price) == ("Сахар 1 кг", 70)

    def test_unchanged_price_is_not_revalued(self) -> None:
        _, sugar, _ = self.known_supplier()

        assert self.edit(
            product_id=sugar.id,
            name="Сахар 1 кг",
            retail_price=65,
            quantity="4",
            expected=422,
        ) == {"detail": "Цена уже 65 ₽, переоценивать нечего"}

        assert (sugar.name, self.store.entries) == ("Сахар-песок 1 кг", [])

    def test_unknown_product_and_closed_period(self) -> None:
        _, sugar, _ = self.known_supplier()
        self.edit(
            product_id=999,
            name="Сахар",
            retail_price=70,
            expected=404,
        )
        self.stocktake(
            counted="0",
        )

        self.edit(
            product_id=sugar.id,
            name="Сахар-песок 1 кг",
            retail_price=70,
            quantity="1",
            shift=-1,
            expected=409,
        )

        assert sugar.retail_price == 65

    @pytest.mark.parametrize(
        argnames=("retail_price", "quantity", "has_price_tag"),
        argvalues=[
            # only the name is corrected, the tag shows the same price
            (65, None, True),
            (70, None, False),
            (70, "4", False),
        ],
    )
    def test_a_new_price_resets_the_price_tag(
        self,
        retail_price: int,
        quantity: str | None,
        has_price_tag: bool,
    ) -> None:
        _, sugar, _ = self.known_supplier()
        sugar.has_price_tag = True

        self.edit(
            product_id=sugar.id,
            name="Сахар 1 кг",
            retail_price=retail_price,
            quantity=quantity,
        )

        assert sugar.has_price_tag is has_price_tag


class TestPriceTag(LedgerApi):
    def mark(
        self,
        product_id: int,
        has_price_tag: bool,
        retail_price: int,
        expected: int = 204,
    ) -> Any:
        response = self.client.put(
            url=f"/product/{product_id}/price-tag",
            json={
                "has_price_tag": has_price_tag,
                "retail_price": retail_price,
            },
        )
        assert response.status_code == expected, response.text
        return response.json() if response.content else None

    def test_marks_and_unmarks_the_tag(self) -> None:
        _, sugar, _ = self.known_supplier()

        self.mark(
            product_id=sugar.id,
            has_price_tag=True,
            retail_price=65,
        )
        assert sugar.has_price_tag is True
        catalog = self.client.get(
            url="/catalog",
        ).json()
        assert [product["has_price_tag"] for product in catalog["products"]] == [
            True,
            False,
        ]

        self.mark(
            product_id=sugar.id,
            has_price_tag=False,
            retail_price=65,
        )
        assert sugar.has_price_tag is False

    def test_a_tag_for_an_older_price_is_rejected(self) -> None:
        _, sugar, _ = self.known_supplier()

        assert self.mark(
            product_id=sugar.id,
            has_price_tag=True,
            retail_price=60,
            expected=409,
        ) == {"detail": "Цена товара изменилась, обновите страницу"}

        assert sugar.has_price_tag is False
        self.mark(
            product_id=999,
            has_price_tag=True,
            retail_price=60,
            expected=404,
        )

    def test_revaluation_and_its_reversal_reset_the_tag(self) -> None:
        _, sugar, salt = self.known_supplier()
        sugar.has_price_tag = True
        salt.has_price_tag = True

        posted = self.revalue(
            lines=[
                (sugar.id, "4", 70),
            ],
        )

        assert (sugar.has_price_tag, salt.has_price_tag) == (False, True)
        sugar.has_price_tag = True
        self.reverse(
            entry_id=posted["id"],
        )
        # the old price is back on the shelf, the tag hung for the new one is wrong
        assert (sugar.retail_price, sugar.has_price_tag) == (65, False)


class TestRevaluationLines(LedgerApi):
    def lines(
        self,
        entry_id: int,
        expected: int = 200,
    ) -> Any:
        response = self.client.get(
            url=f"/ledger/{entry_id}/revaluation",
        )
        assert response.status_code == expected, response.text
        return response.json()

    def test_shows_what_was_revalued_and_what_the_reversal_takes_back(self) -> None:
        _, sugar, salt = self.known_supplier()
        posted = self.revalue(
            lines=[
                (sugar.id, "4", 60),
                (salt.id, "2.5", 40),
            ],
        )
        reversal = self.reverse(
            entry_id=posted["id"],
        )

        assert self.lines(
            entry_id=posted["id"],
        ) == [
            {
                "product_id": sugar.id,
                "product_name": "Сахар-песок 1 кг",
                "quantity": "4",
                "old_price": "65",
                "new_price": "60",
                "amount": "-20.00",
            },
            {
                "product_id": salt.id,
                "product_name": "Соль",
                "quantity": "2.5",
                "old_price": "36",
                "new_price": "40",
                "amount": "10.00",
            },
        ]
        assert [
            line["amount"]
            for line in self.lines(
                entry_id=reversal["id"],
            )
        ] == ["20.00", "-10.00"]

    def test_other_entries_have_no_lines(self) -> None:
        sale = self.sale(
            amount="10.00",
        )

        self.lines(
            entry_id=sale["id"],
            expected=404,
        )
        self.lines(
            entry_id=999,
            expected=404,
        )


class TestReport(LedgerApi):
    def report(
        self,
        date_from: str,
        date_to: str,
        url: str = "/ledger/report",
    ) -> Any:
        return self.client.get(
            url=url,
            params={
                "date_from": date_from,
                "date_to": date_to,
            },
        )

    def make_ledger(self) -> None:
        """a week ago before the period, the rest inside it"""
        self.stock(
            amount="5000.00",
            shift=-7,
        )
        self.stock(
            amount="1000.00",
            shift=-3,
        )
        sale = self.sale(
            amount="700.00",
            shift=-2,
        )
        self.sale(
            amount="300.00",
            cash="100.00",
            shift=-2,
        )
        self.reverse(
            entry_id=sale["id"],
        )
        self.post(
            url="/ledger/write-off",
            payload={
                "occurred_on": day(
                    shift=-1,
                ),
                "amount": "50.00",
                "comment": None,
            },
        )
        _, sugar, _ = self.known_supplier()
        self.revalue(
            lines=[
                (sugar.id, "2", 70),
            ],
            shift=-1,
        )
        self.stocktake(
            counted="5600.00",
        )
        # after the stocktake on the same day
        self.stock(
            amount="100.00",
            shift=0,
        )
        self.cash(
            kind="cash-in",
            amount="500.00",
        )
        self.cash(
            kind="cash-out",
            amount="40.00",
            comment="пакеты",
        )

    def test_totals_add_up_from_the_opening_to_the_closing(self) -> None:
        self.make_ledger()

        report = self.report(
            date_from=day(
                shift=-3,
            ),
            date_to=day(),
        ).json()

        assert {
            key: report[key]
            for key in (
                "opening",
                "invoices",
                "sales",
                "write_offs",
                "stocktakes",
                "revaluations",
                "closing",
            )
        } == {
            "opening": "5000.00",
            # the inflow of today is in the period too
            "invoices": "1100.00",
            # the reversed sale cancels out
            "sales": "-300.00",
            "write_offs": "-50.00",
            "stocktakes": "-60.00",
            "revaluations": "10.00",
            "closing": "5700.00",
        }
        assert {
            key: report[key]
            for key in (
                "cash_opening",
                "sales_cash",
                "sales_card",
                "cash_ins",
                "cash_outs",
                "cash_counts",
                "cash_closing",
            )
        } == {
            "cash_opening": "0",
            "sales_cash": "100.00",
            "sales_card": "200.00",
            "cash_ins": "500.00",
            "cash_outs": "-40.00",
            "cash_counts": "0",
            "cash_closing": "560.00",
        }
        assert [
            (row["kind"], row["amount"], row["balance_after"], row["cash_after"])
            for row in report["entries"]
        ] == [
            ("invoice", "1000.00", "6000.00", "0"),
            ("sale", "-700.00", "5300.00", "0.00"),
            ("sale", "-300.00", "5000.00", "100.00"),
            ("sale", "700.00", "5700.00", "100.00"),
            ("write_off", "-50.00", "5650.00", "100.00"),
            ("revaluation", "10.00", "5660.00", "100.00"),
            ("stocktake", "-60.00", "5600.00", "100.00"),
            ("invoice", "100.00", "5700.00", "100.00"),
            ("cash_in", "0", "5700.00", "600.00"),
            ("cash_out", "0", "5700.00", "560.00"),
        ]

    def test_period_bounds_are_inclusive_and_ordered(self) -> None:
        self.make_ledger()

        report = self.report(
            date_from=day(
                shift=-2,
            ),
            date_to=day(
                shift=-2,
            ),
        ).json()
        backwards = self.report(
            date_from=day(),
            date_to=day(
                shift=-1,
            ),
        )

        assert (report["opening"], report["sales"], report["closing"]) == (
            "6000.00",
            "-300.00",
            "5700.00",
        )
        assert (backwards.status_code, backwards.json()["detail"]) == (
            422,
            "Начало периода позже его конца",
        )

    def test_xlsx_has_the_totals_and_the_entries(self) -> None:
        self.make_ledger()

        response = self.report(
            date_from=day(
                shift=-3,
            ),
            date_to=day(),
            url="/ledger/report.xlsx",
        )

        assert response.status_code == 200, response.text
        assert (
            response.headers["content-disposition"]
            == f'attachment; filename="otchet_{
                day(
                    shift=-3,
                )
            }_{day()}.xlsx"'
        )
        workbook = load_workbook(
            filename=BytesIO(
                initial_bytes=response.content,
            ),
        )
        assert workbook.sheetnames == ["Итоги", "Записи"]
        totals = {
            label: amount
            for label, amount in workbook["Итоги"].iter_rows(
                min_row=3,
                values_only=True,
            )
        }
        assert (totals["Остаток на начало"], totals["Остаток на конец"]) == (5000, 5700)
        assert (
            totals["Касса на начало"],
            totals["Продажи наличными"],
            totals["Касса на конец"],
            totals["Продажи безналом"],
        ) == (0, 100, 560, 200)
        rows = list(
            workbook["Записи"].iter_rows(
                min_row=2,
                values_only=True,
            ),
        )
        assert [(row[2], row[5]) for row in rows[1:4]] == [
            (f"Продажа, сторнирована записью №{rows[3][0]}", -700),
            ("Продажа", -300),
            (f"Продажа, сторно №{rows[1][0]}", 700),
        ]


class TestPeriodReport(LedgerApi):
    def make_ledger(self) -> int:
        """a stock before the period, then sales, a cash-in and a confirmed invoice paid in part"""
        self.stock(
            amount="5000.00",
            shift=-7,
        )
        sale = self.sale(
            amount="700.00",
            cash="700.00",
            shift=-2,
        )
        self.sale(
            amount="300.00",
            cash="100.00",
            shift=-2,
        )
        self.reverse(
            entry_id=sale["id"],
        )
        self.cash(
            kind="cash-in",
            amount="50.00",
            comment="размен",
            shift=-1,
        )
        self.known_supplier()
        draft = self.upload()
        self.post(
            url=f"/invoice/{draft['id']}/confirm",
            payload={
                "occurred_on": day(),
            },
        )
        self.post(
            url=f"/invoice/{draft['id']}/payments",
            payload={
                "paid_on": day(),
                "amount": "100.00",
                "comment": None,
            },
        )
        return draft["id"]

    def days(
        self,
        date_from: str,
        date_to: str,
    ) -> dict[str, Any]:
        response = self.client.get(
            url="/ledger/report/days",
            params={
                "date_from": date_from,
                "date_to": date_to,
            },
        )
        assert response.status_code == 200, response.text
        return response.json()

    def test_period_by_days(self) -> None:
        self.make_ledger()

        report = self.days(
            date_from=day(
                shift=-3,
            ),
            date_to=day(),
        )

        sums = (
            "sales_cash",
            "sales_card",
            "cash_moves",
            "cash_after",
            "balance_after",
        )
        assert [
            Decimal(
                value=report[key],
            )
            for key in ("opening", "cash_opening")
        ] == [5000, 0]
        assert [
            (
                row["day"],
                *[
                    Decimal(
                        value=row[key],
                    )
                    for key in sums
                ],
            )
            for row in report["days"]
        ] == [
            (day(shift=-2), 100, 200, 0, 100, 4700),
            (day(shift=-1), 0, 0, 50, 150, 4700),
            # the invoice is on the day it was posted
            (
                day(),
                0,
                0,
                0,
                150,
                Decimal(
                    value="4898.50",
                ),
            ),
        ]
        assert report["days"][1]["cash_notes"] == [
            {"kind": "cash_in", "cash": "50.00", "comment": "размен"},
        ]
        assert report["days"][2]["invoices"] == [
            {
                "supplier_name": "ООО «Поставщик»",
                "invoice_number": "00042/А",
                "cost_total": "155.00",
                "retail_total": "198.50",
                "paid": "100.00",
            },
        ]
        assert report["paid_in_period"] == "100.00"
        assert [(debt["supplier_name"], debt["debt"]) for debt in report["debts"]] == [
            ("ООО «Поставщик»", "55.00"),
        ]

    def test_reports_print_the_short_name(self) -> None:
        self.make_ledger()
        [supplier] = self.store.suppliers.values()
        response = self.client.put(
            url=f"/supplier/{supplier.id}",
            json={
                "name": supplier.name,
                "short_name": "Поставщик",
                "inn": supplier.inn,
            },
        )
        assert response.status_code == 200, response.text

        days = self.days(
            date_from=day(),
            date_to=day(),
        )
        report = self.client.get(
            url="/ledger/report",
            params={
                "date_from": day(),
                "date_to": day(),
            },
        ).json()
        ledger = self.client.get(
            url="/ledger",
        ).json()

        assert [row["supplier_name"] for row in days["days"][0]["invoices"]] == [
            "Поставщик",
        ]
        assert [debt["supplier_name"] for debt in days["debts"]] == ["Поставщик"]
        assert [
            entry["supplier_name"]
            for entry in report["entries"]
            if entry["kind"] == "invoice"
        ] == ["Поставщик"]
        # the ledger keeps the full name, the stock before the period has no invoice
        assert [
            entry["supplier_name"]
            for entry in ledger["entries"]
            if entry["invoice_id"] is not None
        ] == ["ООО «Поставщик»"]

    def test_debt_is_counted_on_the_last_day(self) -> None:
        self.make_ledger()

        report = self.days(
            date_from=day(
                shift=-3,
            ),
            date_to=day(
                shift=-1,
            ),
        )

        # the invoice was posted today, the day before nothing was owed
        assert report["debts"] == []
        assert (
            Decimal(
                value=report["paid_in_period"],
            )
            == 0
        )

    def test_pdf(self) -> None:
        self.make_ledger()

        response = self.client.get(
            url="/ledger/report.pdf",
            params={
                "date_from": day(
                    shift=-3,
                ),
                "date_to": day(),
            },
        )

        assert response.status_code == 200, response.text
        assert response.headers["content-type"] == "application/pdf"
        assert (
            response.headers["content-disposition"]
            == f'attachment; filename="otchet_{day(shift=-3)}_{day()}.pdf"'
        )
        assert response.content.startswith(
            b"%PDF",
        )

    def test_pdf_of_an_empty_period(self) -> None:
        response = self.client.get(
            url="/ledger/report.pdf",
            params={
                "date_from": day(),
                "date_to": day(),
            },
        )

        assert response.status_code == 200, response.text
        assert response.content.startswith(
            b"%PDF",
        )

    def test_reversed_period_is_rejected(self) -> None:
        response = self.client.get(
            url="/ledger/report.pdf",
            params={
                "date_from": day(),
                "date_to": day(
                    shift=-1,
                ),
            },
        )

        assert response.status_code == 422


class TestHistory(LedgerApi):
    def test_latest_first_with_running_balance_and_pages(self) -> None:
        self.stock(
            amount="1000.00",
            shift=-2,
        )
        self.sale(
            amount="100.00",
        )
        # entered later but dated earlier, it goes before the sale of today
        self.sale(
            amount="50.00",
            shift=-1,
        )

        page = self.history()
        first = self.client.get(
            url="/ledger",
            params={
                "limit": 2,
            },
        ).json()

        assert [
            (row["occurred_on"], row["amount"], row["balance_after"])
            for row in page["entries"]
        ] == [
            (day(), "-100.00", "850.00"),
            (
                day(
                    shift=-1,
                ),
                "-50.00",
                "950.00",
            ),
            (
                day(
                    shift=-2,
                ),
                "1000.00",
                "1000.00",
            ),
        ]
        assert (page["balance"], page["has_more"]) == ("850.00", False)
        assert (
            len(
                first["entries"],
            ),
            first["has_more"],
        ) == (2, True)
