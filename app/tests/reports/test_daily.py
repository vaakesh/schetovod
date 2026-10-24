from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from tests.ledger.helpers import LedgerApi, day


class TestReport(LedgerApi):
    def report(
        self,
        date_from: str,
        date_to: str,
        url: str = "/ledger/report",
    ) -> Any:
        return self.client.get(
            url=url,
            params={
                "date_from": date_from,
                "date_to": date_to,
            },
        )

    def make_ledger(self) -> None:
        """a week ago before the period, the rest inside it"""
        self.stock(
            amount="5000.00",
            shift=-7,
        )
        self.stock(
            amount="1000.00",
            shift=-3,
        )
        sale = self.sale(
            amount="700.00",
            shift=-2,
        )
        self.sale(
            amount="300.00",
            cash="100.00",
            shift=-2,
        )
        self.reverse(
            entry_id=sale["id"],
        )
        self.post(
            url="/ledger/write-off",
            payload={
                "occurred_on": day(
                    shift=-1,
                ),
                "amount": "50.00",
                "comment": None,
            },
        )
        _, sugar, _ = self.known_supplier()
        self.revalue(
            lines=[
                (sugar.id, "2", 70),
            ],
            shift=-1,
        )
        self.stocktake(
            counted="5600.00",
        )
        # after the stocktake on the same day
        self.stock(
            amount="100.00",
            shift=0,
        )
        self.cash(
            kind="cash-in",
            amount="500.00",
        )
        self.cash(
            kind="cash-out",
            amount="40.00",
            comment="пакеты",
        )

    def test_totals_add_up_from_the_opening_to_the_closing(self) -> None:
        self.make_ledger()

        report = self.report(
            date_from=day(
                shift=-3,
            ),
            date_to=day(),
        ).json()

        assert {
            key: report[key]
            for key in (
                "opening",
                "invoices",
                "sales",
                "write_offs",
                "stocktakes",
                "revaluations",
                "closing",
            )
        } == {
            "opening": "5000.00",
            # the inflow of today is in the period too
            "invoices": "1100.00",
            # the reversed sale cancels out
            "sales": "-300.00",
            "write_offs": "-50.00",
            "stocktakes": "-60.00",
            "revaluations": "10.00",
            "closing": "5700.00",
        }
        assert {
            key: report[key]
            for key in (
                "cash_opening",
                "sales_cash",
                "sales_card",
                "cash_ins",
                "cash_outs",
                "cash_counts",
                "cash_closing",
            )
        } == {
            "cash_opening": "0",
            "sales_cash": "100.00",
            "sales_card": "200.00",
            "cash_ins": "500.00",
            "cash_outs": "-40.00",
            "cash_counts": "0",
            "cash_closing": "560.00",
        }
        assert [
            (row["kind"], row["amount"], row["balance_after"], row["cash_after"])
            for row in report["entries"]
        ] == [
            ("invoice", "1000.00", "6000.00", "0"),
            ("sale", "-700.00", "5300.00", "0.00"),
            ("sale", "-300.00", "5000.00", "100.00"),
            ("sale", "700.00", "5700.00", "100.00"),
            ("write_off", "-50.00", "5650.00", "100.00"),
            ("revaluation", "10.00", "5660.00", "100.00"),
            ("stocktake", "-60.00", "5600.00", "100.00"),
            ("invoice", "100.00", "5700.00", "100.00"),
            ("cash_in", "0", "5700.00", "600.00"),
            ("cash_out", "0", "5700.00", "560.00"),
        ]

    def test_period_bounds_are_inclusive_and_ordered(self) -> None:
        self.make_ledger()

        report = self.report(
            date_from=day(
                shift=-2,
            ),
            date_to=day(
                shift=-2,
            ),
        ).json()
        backwards = self.report(
            date_from=day(),
            date_to=day(
                shift=-1,
            ),
        )

        assert (report["opening"], report["sales"], report["closing"]) == (
            "6000.00",
            "-300.00",
            "5700.00",
        )
        assert (backwards.status_code, backwards.json()["detail"]) == (
            422,
            "Начало периода позже его конца",
        )

    def test_xlsx_has_the_totals_and_the_entries(self) -> None:
        self.make_ledger()

        response = self.report(
            date_from=day(
                shift=-3,
            ),
            date_to=day(),
            url="/ledger/report.xlsx",
        )

        assert response.status_code == 200, response.text
        assert (
            response.headers["content-disposition"]
            == f'attachment; filename="otchet_{
                day(
                    shift=-3,
                )
            }_{day()}.xlsx"'
        )
        workbook = load_workbook(
            filename=BytesIO(
                initial_bytes=response.content,
            ),
        )
        assert workbook.sheetnames == ["Итоги", "Записи"]
        totals = {
            label: amount
            for label, amount in workbook["Итоги"].iter_rows(
                min_row=3,
                values_only=True,
            )
        }
        assert (totals["Остаток на начало"], totals["Остаток на конец"]) == (5000, 5700)
        assert (
            totals["Касса на начало"],
            totals["Продажи наличными"],
            totals["Касса на конец"],
            totals["Продажи безналом"],
        ) == (0, 100, 560, 200)
        rows = list(
            workbook["Записи"].iter_rows(
                min_row=2,
                values_only=True,
            ),
        )
        assert [(row[2], row[5]) for row in rows[1:4]] == [
            (f"Продажа, сторнирована записью №{rows[3][0]}", -700),
            ("Продажа", -300),
            (f"Продажа, сторно №{rows[1][0]}", 700),
        ]
