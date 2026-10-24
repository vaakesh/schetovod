from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from core.enums import LedgerKind
from core.schema_types import format_invoice_date
from reports.schemas import LedgerReport

MONEY_FORMAT = "#,##0.00"
DATE_FORMAT = "DD.MM.YYYY"

# the same labels as the ledger page
KIND_LABELS = {
    LedgerKind.INVOICE: "Накладная",
    LedgerKind.SALE: "Продажа",
    LedgerKind.WRITE_OFF: "Списание",
    LedgerKind.STOCKTAKE: "Переучёт",
    LedgerKind.REVALUATION: "Переоценка",
    LedgerKind.CASH_IN: "Внесение",
    LedgerKind.CASH_OUT: "Выдача",
    LedgerKind.CASH_COUNT: "Сверка кассы",
}


def report_filename(report: LedgerReport) -> str:
    date_from = format_invoice_date(
        value=report.date_from,
    )
    date_to = format_invoice_date(
        value=report.date_to,
    )
    return f"otchet_{date_from}_{date_to}.xlsx"


def report_workbook(report: LedgerReport) -> bytes:
    """the totals on the first sheet, the entries of the period on the second"""
    workbook = Workbook()
    bold = Font(
        bold=True,
    )

    totals = workbook.active
    if totals is None:
        raise RuntimeError(
            "a new workbook has a sheet",
        )
    totals.title = "Итоги"
    date_from = format_invoice_date(
        value=report.date_from,
    )
    date_to = format_invoice_date(
        value=report.date_to,
    )
    totals.append(
        [
            f"Отчёт с {date_from} по {date_to}",
        ],
    )
    totals["A1"].font = bold
    totals.append(
        [],
    )
    # the stock and the till, each block adds up from its first line to its last
    blocks = (
        (
            ("Остаток на начало", report.opening),
            ("Приход по накладным", report.invoices),
            ("Продажи", report.sales),
            ("Списания", report.write_offs),
            ("Переучёты", report.stocktakes),
            ("Переоценка", report.revaluations),
            ("Остаток на конец", report.closing),
        ),
        (
            ("Касса на начало", report.cash_opening),
            ("Продажи наличными", report.sales_cash),
            ("Внесения", report.cash_ins),
            ("Выдачи", report.cash_outs),
            ("Сверки кассы", report.cash_counts),
            ("Касса на конец", report.cash_closing),
        ),
    )
    for index, block in enumerate(
        blocks,
    ):
        if index:
            totals.append(
                [],
            )
        first = totals.max_row + 1
        for label, amount in block:
            totals.append(
                [
                    label,
                    amount,
                ],
            )
            totals.cell(
                row=totals.max_row,
                column=2,
            ).number_format = MONEY_FORMAT
        for row in (first, totals.max_row):
            for column in (1, 2):
                totals.cell(
                    row=row,
                    column=column,
                ).font = bold
    totals.append(
        [
            "Продажи безналом",
            report.sales_card,
        ],
    )
    totals.cell(
        row=totals.max_row,
        column=2,
    ).number_format = MONEY_FORMAT
    totals.column_dimensions["A"].width = 24
    totals.column_dimensions["B"].width = 16

    entries = workbook.create_sheet(
        title="Записи",
    )
    entries.append(
        [
            "№",
            "Дата",
            "Операция",
            "Накладная",
            "Поставщик",
            "Сумма",
            "Остаток",
            "Касса",
            "Остаток кассы",
            "Комментарий",
        ],
    )
    for cell in entries[1]:
        cell.font = bold
    for entry in report.entries:
        operation = KIND_LABELS[entry.kind]
        if entry.reverses_id is not None:
            operation = f"{operation}, сторно №{entry.reverses_id}"
        elif entry.reversed_by_id is not None:
            operation = f"{operation}, сторнирована записью №{entry.reversed_by_id}"
        entries.append(
            [
                entry.id,
                entry.occurred_on,
                operation,
                entry.invoice_number,
                entry.supplier_name,
                entry.amount,
                entry.balance_after,
                entry.cash,
                entry.cash_after,
                entry.comment,
            ],
        )
        row = entries.max_row
        entries.cell(
            row=row,
            column=2,
        ).number_format = DATE_FORMAT
        for column in (6, 7, 8, 9):
            entries.cell(
                row=row,
                column=column,
            ).number_format = MONEY_FORMAT
    for column, width in zip(
        "ABCDEFGHIJ",
        (8, 12, 26, 14, 28, 14, 14, 14, 14, 40),
        strict=True,
    ):
        entries.column_dimensions[column].width = width
    entries.freeze_panes = "A2"

    buffer = BytesIO()
    workbook.save(
        filename=buffer,
    )
    return buffer.getvalue()
